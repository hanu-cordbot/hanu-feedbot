import asyncio
import json
import random
import sys
import time

import discord
from pathlib import Path

sys.path.append(str(Path(__file__).resolve().parents[1]))

from studybot import ingest, questions  # noqa: E402
from studybot.db import Store  # noqa: E402

PASSAGE = (
    "Enzymes speed up reactions by lowering the activation energy required for the reaction to proceed. "
    "They do not change the equilibrium position of a reaction. Competitive inhibitors bind to the active site "
    "and can be overcome by increasing substrate concentration, whereas non-competitive inhibitors bind elsewhere "
    "and reduce the maximum rate regardless of how much substrate is added. Temperature increases reaction rate "
    "until the enzyme denatures, after which activity falls sharply."
)


def draft(**overrides):
    q = {
        "stem": "A drug binds an enzyme's active site. What happens if you greatly increase substrate concentration?",
        "correct_answer": "The inhibition is largely overcome",
        "distractors": [
            {"text": "The maximum rate falls permanently", "misconception": "Confuses competitive with non-competitive"},
            {"text": "The equilibrium shifts toward products", "misconception": "Thinks enzymes move equilibrium"},
            {"text": "The enzyme denatures more quickly", "misconception": "Mixes up substrate and temperature effects"},
        ],
        "explanation": "Competitive inhibitors compete for the active site, so more substrate outcompetes them.",
        "source_quote": "Competitive inhibitors bind to the active site and can be overcome by increasing substrate concentration",
        "difficulty": "medium",
        "cognitive_level": "apply",
    }
    q.update(overrides)
    return q


def test_chunking_keeps_paragraphs_and_overlaps():
    text = "\n\n".join(f"Paragraph {i}. " + "word " * 120 for i in range(20))
    chunks = ingest.chunk_text(text, size=1500, overlap=200)
    assert len(chunks) > 3
    assert all(len(c) <= 1500 + 200 + 10 for c in chunks)
    assert all(f"Paragraph {i}." in text for i in range(20))
    joined = " ".join(chunks)
    assert all(f"Paragraph {i}." in joined for i in range(20))


def test_extract_plain_text_and_rejects_unknown():
    assert ingest.extract_text("notes.md", b"Hello\r\n\r\n\r\nworld") == "Hello\n\nworld"
    try:
        ingest.extract_text("song.mp3", b"x")
    except ingest.UnsupportedFile:
        pass
    else:
        raise AssertionError("expected UnsupportedFile")


def test_local_checks_accept_good_and_catch_cues():
    assert questions.local_problems(draft(), PASSAGE, []) is None
    assert "all/none" in questions.local_problems(
        draft(distractors=[*draft()["distractors"][:2], {"text": "All of the above", "misconception": "x"}]), PASSAGE, []
    )
    assert "longest" in questions.local_problems(
        draft(correct_answer="The inhibition is largely overcome because substrate molecules outcompete the inhibitor for the site"),
        PASSAGE,
        [],
    )
    assert "quote" in questions.local_problems(draft(source_quote="Enzymes are made of DNA and act as genes in the nucleus"), PASSAGE, [])
    assert "duplicate" in questions.local_problems(draft(), PASSAGE, [draft()["stem"]])
    assert "refers to the source" in questions.local_problems(draft(stem="According to the passage, what do enzymes do?"), PASSAGE, [])


def test_shuffle_tracks_correct_answer():
    for seed in range(10):
        q = questions.shuffle_options(draft(), random.Random(seed))
        assert q["options"][q["correct_index"]] == "The inhibition is largely overcome"
        assert q["rationales"][q["correct_index"]] == ""
        assert sorted(q["options"]) == sorted([draft()["correct_answer"]] + [d["text"] for d in draft()["distractors"]])


def test_pipeline_drops_questions_the_blind_solver_gets_wrong(monkeypatch):
    calls = []

    async def fake_generate_json(prompt, *, system, audio=None, temperature=0.7, tries=3):
        calls.append(system)
        if system == questions.WRITER_SYSTEM:
            return {"questions": [draft(), draft(stem="Which change would raise an enzyme's activity until it suddenly collapses?",
                                              correct_answer="Raising the temperature",
                                              source_quote="Temperature increases reaction rate until the enzyme denatures")]}
        if system == questions.SOLVER_SYSTEM:
            items = json.loads(prompt.split("QUESTIONS:\n", 1)[1].split("\n\nFor each", 1)[0])
            answers = []
            for item in items:
                letter = next(k for k, v in item["options"].items() if v in ("The inhibition is largely overcome", "Raising the temperature"))
                # Solver agrees on the first question and disagrees on the second.
                if item["id"] == 1:
                    letter = "ABCD"[("ABCD".index(letter) + 1) % 4]
                answers.append({"id": item["id"], "choice": letter, "ambiguous": False, "reason": ""})
            return {"answers": answers}
        return {"reviews": [{"id": 0, "verdict": "keep", "problems": "",
                             "scores": {k: 5 for k in ("grounded", "unambiguous", "distractors", "learning_value", "standalone")}}]}

    monkeypatch.setattr(questions.llm, "generate_json", fake_generate_json)
    report = asyncio.run(questions.generate_from_passage(PASSAGE, 2, [], rng=random.Random(1)))
    assert calls == [questions.WRITER_SYSTEM, questions.SOLVER_SYSTEM, questions.REVIEWER_SYSTEM]
    assert len(report.accepted) == 1 and report.accepted[0]["review_score"] == 5
    assert any("blind solver disagreed" in reason for _, reason in report.rejected)


def test_store_roundtrip_ordering_and_flags():
    store = Store(":memory:")
    deck = store.get_or_create_deck(1, "Bio")
    assert store.get_or_create_deck(1, " bio ") == deck
    doc = store.add_document(deck, "notes.txt", "abc", [PASSAGE], user_id=7)
    assert store.add_document(deck, "notes.txt", "abc", [PASSAGE], user_id=7) is None
    chunk = store.pick_chunks(deck, 1, focus="inhibitors")[0]
    assert chunk["document_id"] == doc

    ids = [store.add_question(deck, chunk["id"], questions.shuffle_options(draft(stem=f"Question number {i} about enzymes?"), random.Random(i)))
           for i in range(3)]
    first = store.next_questions(deck, user_id=42, count=3)
    assert {q["id"] for q in first} == set(ids)

    store.record_attempt(ids[0], 42, 0, correct=False)
    store.record_attempt(ids[1], 42, 0, correct=True)
    order = [q["id"] for q in store.next_questions(deck, 42, 3)]
    assert order == [ids[2], ids[0], ids[1]]  # unseen, then missed, then known
    assert [q["id"] for q in store.next_questions(deck, 42, 3, skip_recent_seconds=600)] == [ids[2]]

    assert store.flag_question(ids[2], 1) is False
    assert store.flag_question(ids[2], 1) is False  # same person twice counts once
    assert store.flag_question(ids[2], 2) is True
    assert ids[2] not in [q["id"] for q in store.next_questions(deck, 42, 3)]

    stats = store.user_stats(42, deck)
    assert (stats["answered"], stats["correct"]) == (2, 1)


def test_login_budget(tmp_path, monkeypatch):
    from studybot import __main__ as entry, config

    monkeypatch.setattr(config, "DATA_DIR", tmp_path)
    monkeypatch.setattr(config, "MAX_LOGINS_PER_DAY", 3)
    assert [entry.take_login_slot() for _ in range(3)] == [0, 0, 0]
    assert entry.take_login_slot() > 23 * 3600
    stale = [time.time() - 25 * 3600] * 3
    (tmp_path / "logins.json").write_text(json.dumps(stale))
    assert entry.take_login_slot() == 0


def _deck_with_questions(n=3):
    store = Store(":memory:")
    deck = store.get_or_create_deck(9, "Enzymes 101")
    store.add_document(deck, "bio-notes.pdf", "h", [PASSAGE], user_id=1)
    for i in range(n):
        q = draft(stem=f"Question {i}: what happens to inhibition when substrate rises?")
        if i == 0:
            q["explanation"] = "Watch out: </script><script>alert(1)</script> is just text."
        store.add_question(deck, None, questions.shuffle_options(q, random.Random(i)))
    return store, deck


def test_export_html_embeds_questions_safely():
    from studybot import export

    store, deck = _deck_with_questions()
    data = export.pack_data(store, deck, "enzymes 101")
    page = export.render_html(data)
    assert data["sources"] == ["bio-notes.pdf"] and len(data["questions"]) == 3
    assert "</script><script>alert" not in page  # document text can't break out of the data block
    embedded = page.split('type="application/json">', 1)[1].split("</script>", 1)[0]
    assert json.loads(embedded)["questions"][0]["options"] == data["questions"][0]["options"]


def test_export_anki_and_index(tmp_path):
    from studybot import export

    store, _ = _deck_with_questions(2)
    written = export.export_guild(store, 9, tmp_path)
    assert sorted(p.name for p in written) == ["enzymes-101.anki.txt", "enzymes-101.html", "index.html"]
    notes = (tmp_path / "enzymes-101.anki.txt").read_text().splitlines()
    assert notes[:2] == ["#separator:tab", "#html:true"]
    assert all(len(line.split("\t")) == 3 for line in notes[3:])
    assert 'href="enzymes-101.html"' in (tmp_path / "index.html").read_text()


# --- extracting existing questions -------------------------------------------

def test_parser_handles_common_layouts():
    from studybot.extract import parse_mcqs

    text = """Unit 3 Quiz
Choose the best answer.

1. Which organelle produces ATP?
A. Nucleus
B. Mitochondrion
C. Ribosome
D. Golgi body
Answer: B
Explanation: Mitochondria run
oxidative phosphorylation.

Q2) Capital of France? (A) Berlin (B) Paris (C) Rome

Which molecule carries amino acids to the ribosome?
a) mRNA
b) tRNA *
c) rRNA

4. Who wrote Faust?
A. Goethe
B. Schiller

Answer Key
2. B
"""
    qs = parse_mcqs(text)
    assert [q.stem for q in qs] == [
        "Which organelle produces ATP?",
        "Capital of France?",
        "Which molecule carries amino acids to the ribosome?",
        "Who wrote Faust?",
    ]
    assert [q.answer for q in qs] == [1, 1, 1, None]  # answer line, key section, star, none given
    assert qs[0].explanation == "Mitochondria run oxidative phosphorylation."
    assert qs[1].options == ["Berlin", "Paris", "Rome"]
    assert parse_mcqs("Enzymes speed reactions. A. Smith wrote about it. B. Jones agreed.") == []


def test_docx_bold_option_marks_the_answer(tmp_path):
    import docx
    from studybot.extract import parse_mcqs

    d = docx.Document()
    d.add_paragraph("1. Which gas do plants absorb?")
    d.add_paragraph("A. Oxygen")
    d.add_paragraph().add_run("B. Carbon dioxide").bold = True
    d.add_paragraph("C. Nitrogen")
    path = tmp_path / "q.docx"
    d.save(path)
    qs = parse_mcqs(ingest.extract_text("q.docx", path.read_bytes()))
    assert len(qs) == 1 and qs[0].answer == 1 and qs[0].options[1] == "Carbon dioxide"


def test_ai_copy_drops_questions_not_in_the_file(monkeypatch):
    from studybot import extract

    source = "| 1 | What is the boiling point of water at sea level? | 90 C | 100 C | 110 C |\n" * 2

    async def fake(prompt, *, system, audio=None, temperature=0.7, tries=3):
        return {"questions": [
            {"number": 1, "stem": "What is the boiling point of water at sea level?", "options": ["90 C", "100 C", "110 C"], "answer_position": None},
            {"number": 2, "stem": "What is the freezing point of mercury?", "options": ["-39 C", "0 C", "12 C"], "answer_position": 1},
        ]}

    monkeypatch.setattr(extract.llm, "generate_json", fake)
    found, dropped = asyncio.run(extract.copy_with_ai(source))
    assert [q.stem for q in found] == ["What is the boiling point of water at sea level?"] and dropped == 1


def test_find_questions_skips_the_model_when_not_needed(monkeypatch):
    from studybot import extract

    async def boom(*a, **k):
        raise AssertionError("model should not be called")

    monkeypatch.setattr(extract.llm, "generate_json", boom)
    clean = "1. Pick one?\nA. x\nB. y\nAnswer: A\n"
    assert asyncio.run(extract.find_questions("q.txt", clean.encode(), clean)).method == "parser"
    notes = "Plain lecture notes about enzymes. " * 20
    assert asyncio.run(extract.find_questions("n.txt", notes.encode(), notes)).questions == []


def test_ai_answers_are_labelled():
    from studybot import extract

    q = extract.Extracted("Who wrote Hamlet?", ["Marlowe", "Shakespeare"], answer=None)
    assert extract.to_stored(q) is None
    stored = extract.to_stored(q, {"index": 1, "confidence": "high", "reason": "He did."})
    assert stored["answer_source"] == "ai" and stored["correct_index"] == 1
    assert "suggested by AI" in questions.caution(stored)
    assert questions.caution(extract.to_stored(extract.Extracted("Q?", ["a", "b"], answer=0))) == ""
    assert "AI-written" in questions.caution({"origin": "generated"})


def test_old_database_gets_new_columns(tmp_path):
    import sqlite3

    path = tmp_path / "old.sqlite3"
    conn = sqlite3.connect(path)
    conn.execute("CREATE TABLE questions (id INTEGER PRIMARY KEY, deck_id INTEGER, chunk_id INTEGER, stem TEXT, "
                 "options_json TEXT, correct_index INTEGER, explanation TEXT, rationales_json TEXT, source_quote TEXT, "
                 "difficulty TEXT, cognitive_level TEXT, review_score REAL, flags INTEGER DEFAULT 0, "
                 "retired INTEGER DEFAULT 0, created_at REAL)")
    conn.commit()
    conn.close()
    cols = {r["name"] for r in Store(path).conn.execute("PRAGMA table_info(questions)")}
    assert {"origin", "answer_source", "document_id"} <= cols


# --- library, labels, picking, solo runs, live games -------------------------

def _q(origin="extracted", answer_source="document", net=0, flags=0, last=None, qid=1):
    return {"id": qid, "origin": origin, "answer_source": answer_source, "net_rating": net, "flags": flags, "last_correct": last}


def test_labels_and_quality():
    from studybot import quality

    assert quality.labels(_q()) == "📄"
    assert quality.labels(_q(answer_source="ai", net=2, flags=1)) == "🤖 ⭐ 🚩"
    assert quality.labels(_q(origin="generated")) == "✨"
    assert quality.quality(_q()) == 2 and quality.quality(_q(origin="generated")) == -1
    assert quality.quality(_q(net=1, flags=1)) == 1
    assert quality.subject_emoji("Organic Chemistry") == "⚗️" and quality.subject_emoji("Italian") == "📚"


def test_pick_prefers_good_questions_and_tops_up_when_short():
    from studybot import quality

    good = [_q(qid=i) for i in range(10)]
    weak = [_q(origin="generated", net=-1, qid=100 + i) for i in range(10)]
    a = quality.pick(good + weak, 5, random.Random(1))
    assert [q["id"] for q in a] == [q["id"] for q in quality.pick(good + weak, 5, random.Random(1))]  # seeded
    assert all(q["id"] < 100 for q in a)
    b = quality.pick(good[:2] + weak, 5, random.Random(2))
    assert len(b) == 5 and sum(q["id"] < 100 for q in b) == 2
    assert len({q["id"] for q in quality.pick(good, 50, random.Random(3))}) == 10  # never repeats


def test_pick_favours_missed_questions_for_solo():
    from studybot import quality

    missed = [_q(last=0, qid=i) for i in range(5)]
    known = [_q(last=1, qid=10 + i) for i in range(5)]
    rng = random.Random(7)
    hits = sum(q["last_correct"] == 0 for _ in range(400) for q in quality.pick(missed + known, 1, rng, personal=True))
    assert hits > 280  # weights 1.8×2.2 vs 0.6×2.2 → missed should win ~75% of draws


def test_subjects_library_ratings_and_runs(tmp_path):
    import sqlite3

    path = tmp_path / "old.sqlite3"
    conn = sqlite3.connect(path)
    conn.execute("CREATE TABLE decks (id INTEGER PRIMARY KEY, guild_id INTEGER NOT NULL, name TEXT NOT NULL, "
                 "created_by INTEGER, created_at REAL NOT NULL, UNIQUE (guild_id, name))")
    conn.execute("INSERT INTO decks (guild_id, name, created_at) VALUES (1, 'old set', 0)")
    conn.commit()
    conn.close()
    store = Store(path)
    assert store.deck_subject(1) == "general"

    bio = store.get_or_create_deck(1, "cells", subject="Biology")
    assert store.get_or_create_deck(1, "cells") == bio and store.deck_subject(bio) == "biology"
    store.get_or_create_deck(1, "cells", subject="biology 101")
    assert store.subjects(1) == ["biology 101", "general"]

    from studybot import extract

    q1 = store.add_question(bio, None, extract.to_stored(extract.Extracted("Q one?", ["a", "b"], 0)))
    q2 = store.add_question(bio, None, extract.to_stored(extract.Extracted("Q two?", ["a", "b"], None), {"index": 1, "confidence": "low", "reason": ""}))
    assert store.rate_question(q1, 5, 1) == 1 and store.rate_question(q1, 6, 1) == 2
    assert store.rate_question(q1, 6, -1) == 0  # changing your vote replaces it
    store.rate_question(q1, 6, 1)
    row = next(r for r in store.library(1) if r["id"] == bio)
    assert (row["questions"], row["from_file"], row["ai_answer"], row["favourites"]) == (2, 1, 1, 1)
    assert {q["id"]: q["net_rating"] for q in store.question_pool([bio])} == {q1: 2, q2: 0}

    run = store.create_run(9, [q1, q2])
    assert store.advance_run(run, 0, True) and not store.advance_run(run, 0, True)  # double click ignored
    assert store.get_run(run)["position"] == 1 and store.get_run(run)["correct"] == 1


def test_live_scoring():
    from studybot.play import points_for

    assert points_for(False, 10, 20, 3) == 0
    assert points_for(True, 20, 20, 0) == 1000 and points_for(True, 0, 20, 0) == 500
    assert points_for(True, 10, 20, 2) == 850 and points_for(True, 10, 20, 99) == 1000


def test_live_game_is_one_message_with_private_picks():
    import types
    from studybot import play

    store, _ = _deck_with_questions(1)
    qid = store.deck_questions(store.find_deck(9, "enzymes 101"))[0]["id"]
    q = store.get_question(qid)
    edits, sent, deleted = [], [], []

    class Msg:
        async def edit(self, **kw):
            edits.append(kw)

    class Channel:
        async def send(self, *a, **kw):
            assert "attachments" not in kw  # Messageable.send() only takes files=
            sent.append(kw)
            return Msg()

    client = types.SimpleNamespace(store=store, live_games={})

    def interaction(uid, name):
        log = []

        async def send_message(text, **kw):
            log.append((text, kw))

        async def delete_original_response():
            deleted.append(uid)

        return types.SimpleNamespace(
            client=client, user=types.SimpleNamespace(id=uid, display_name=name),
            response=types.SimpleNamespace(send_message=send_message), log=log,
            delete_original_response=delete_original_response,
        )

    async def scenario():
        game = play.LiveGame(1, host_id=1, channel=Channel(), question_ids=[qid], time_limit=30,
                             settings=play.QuizSettings(cards=False))
        client.live_games[1] = game
        await play.ask_next(client, game)
        assert len(sent) == 1 and "🙋 **0** answered" in sent[0]["embed"].fields[0].value
        ids = [getattr(c, "item", c).custom_id for c in sent[0]["view"].children]
        # A/B/C/D are right on the public message, plus the host's controls.
        assert [i.split(":")[1] for i in ids] == ["la"] * len(q["options"]) + ["lpm", "lh", "lh"]

        right, wrong = interaction(2, "Ada"), interaction(3, "Bo")
        await play.LiveAnswer(1, 0, q["correct_index"]).callback(right)
        await play.LiveAnswer(1, 0, (q["correct_index"] + 1) % 4).callback(wrong)
        await play.LiveAnswer(1, 0, 0).callback(right)  # second pick is refused
        text, kw = right.log[0]
        assert kw["ephemeral"] and "Locked in" in text and "✅" in text and "#1 of 1" in text
        assert "already locked in" in right.log[1][0] and right.log[1][1]["delete_after"]
        assert "❌" in wrong.log[0][0]
        assert game.scores[2] > 900 and game.scores[3] == 0
        await asyncio.sleep(play.COUNTER_EDIT_GAP + 0.1)
        assert "🙋 **2** answered" in edits[-1]["embed"].fields[0].value
        assert "answer:" not in edits[-1]["embed"].title  # still hidden before the reveal

        await play.reveal(game, 0)
        revealed = edits[-1]["embed"]
        assert f"answer: {'ABCD'[q['correct_index']]}" in revealed.title
        assert "Ada" in next(f.value for f in revealed.fields if f.name == "Leaderboard")
        assert sorted(deleted) == [2, 3]  # private "locked in" replies are cleaned up at the reveal
        await play.ask_next(client, game)  # past the last question → podium on the same message
        assert len(sent) == 1 and 1 not in client.live_games and "🥇 **Ada**" in edits[-1]["embed"].description
        assert edits[-1]["view"] is None

    asyncio.run(scenario())


def test_bot_commands_and_custom_ids():
    import re as _re
    from studybot import play
    from studybot.bot import StudyBot, StudyCog

    async def build():
        bot = StudyBot(Store(":memory:"))
        await bot.add_cog(StudyCog(bot))
        names = sorted(c.name for c in bot.tree.get_commands())
        for c in bot.tree.get_commands():
            c.to_dict(bot.tree)
        await bot.close()
        return names

    assert asyncio.run(build()) == ["generate", "interview", "library", "quiz", "share", "stats", "upload"]
    big = 10**18
    for item in (play.SoloAnswer(big, 24, 5), play.SoloNext(big), play.LiveAnswer(big, 24, 5),
                 play.LiveHost(big, "reveal"), play.LiveMine(big, 24), play.RateButton(big, False), play.FlagButton(big)):
        cid = item.item.custom_id
        assert len(cid) <= 100 and _re.fullmatch(type(item).__discord_ui_compiled_template__, cid)


def test_codex_provider_routes_text_images_and_audio(monkeypatch):
    from studybot import config, llm

    monkeypatch.setattr(config, "PROVIDER", "codex")
    monkeypatch.setattr(config, "GEMINI_API_KEY", None)
    seen = []

    async def fake_run(prompt, attachment=None):
        seen.append((prompt, attachment))
        return '```json\n{"ok": true}\n```'

    monkeypatch.setattr(llm, "_run_codex", fake_run)
    assert asyncio.run(llm.generate_json("task text", system="be terse")) == {"ok": True}
    assert "be terse" in seen[0][0] and "task text" in seen[0][0] and seen[0][1] is None

    image = (b"\x89PNG", "image/png")
    asyncio.run(llm.generate_json("read it", system="s", audio=image))
    assert seen[1][1] == image

    try:
        asyncio.run(llm.generate_json("grade", system="s", audio=(b"ogg", "audio/ogg")))
    except llm.LLMError as exc:
        assert "GEMINI_API_KEY" in str(exc)
    else:
        raise AssertionError("audio should need Gemini")
    assert len(seen) == 2


def test_codex_lockdown_turns_off_every_tool():
    from studybot import llm

    args = " ".join(llm.LOCKDOWN)
    for feature in ("shell_tool", "unified_exec", "browser_use", "computer_use", "apps", "plugins", "hooks", "memories"):
        assert f"features.{feature}=false" in args
    assert "mcp_servers={}" in args and 'web_search="disabled"' in args


def test_zip_bombs_and_huge_files_are_refused():
    import io as _io
    import zipfile

    buf = _io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr("word/document.xml", b"\0" * (ingest.MAX_UNZIPPED_BYTES + 1))
    try:
        ingest.extract_text("bomb.docx", buf.getvalue())
    except ingest.UnsupportedFile:
        pass
    else:
        raise AssertionError("zip bomb was parsed")
    try:
        ingest.extract_text("broken.pptx", b"not a zip")
    except ingest.UnsupportedFile:
        pass
    else:
        raise AssertionError("broken Office file was parsed")


def test_bot_never_pings_from_uploaded_text():
    from studybot.bot import StudyBot

    async def build():
        bot = StudyBot(Store(":memory:"))
        try:
            m = bot.allowed_mentions
            return m.everyone, m.roles, m.users
        finally:
            await bot.close()

    assert asyncio.run(build()) == (False, False, False)


def test_json_question_banks_are_read_without_ai():
    from studybot import extract

    bank = [
        {"title": "Week 1", "mcqs": [
            {"q": "Capital of France?", "options": ["Berlin", "Paris", "Rome"], "correctIndex": 1, "explanation": "It is Paris."},
            {"question": "2 + 2?", "choices": {"A": "3", "B": "4"}, "answer": "B"},
            {"stem": "Largest planet?", "options": [{"text": "Mars"}, {"text": "Jupiter", "correct": True}]},
            {"q": "Ambiguous answer number", "options": ["x", "y", "z"], "answer": 2},
        ]},
        {"notes": "not a question", "options": "nope"},
    ]
    data = json.dumps(bank).encode()
    found = asyncio.run(extract.find_questions("bank.json", data, ingest.extract_text("bank.json", data)))
    assert found.method == "parser"
    assert [(q.stem, q.answer) for q in found.questions] == [
        ("Capital of France?", 1), ("2 + 2?", 1), ("Largest planet?", 1), ("Ambiguous answer number", None)
    ]
    assert found.questions[0].explanation == "It is Paris."
    assert "Capital of France?" in ingest.extract_text("bank.json", data)
    assert ingest.extract_text("broken.json", b"{not json") == "{not json"


def test_quiz_settings_selection_and_shuffle():
    from studybot import play

    s = play.QuizSettings.from_dict({"length": 3, "order": "bogus", "sources": "file", "unknown": 1})
    assert s.order == "smart" and s.sources == "file" and s.length == 3
    assert play.QuizSettings.from_dict(s.to_dict()) == s and play.QuizSettings(timer=0).live_timer == play.LIVE_DEFAULT_TIMER

    def q(i, origin, source):
        return {"id": i, "origin": origin, "answer_source": source, "net_rating": 0, "flags": 0}

    pool = [q(1, "extracted", "document"), q(2, "extracted", "ai"), q(3, "generated", "ai"), q(4, "extracted", "document")]
    assert [x["id"] for x in play.filter_sources(pool, "file")] == [1, 4]
    assert [x["id"] for x in play.filter_sources(pool, "no_ai_written")] == [1, 2, 4]
    rng = random.Random(3)
    assert [x["id"] for x in play.choose_questions(pool, play.QuizSettings(length=0, order="file"), rng, False)] == [1, 2, 3, 4]
    assert len(play.choose_questions(pool, play.QuizSettings(length=2, order="random"), rng, False)) == 2
    assert len(play.choose_questions(pool, play.QuizSettings(length=50), rng, False)) == 4

    assert play.option_order(4, False, "x") == [0, 1, 2, 3]
    shuffled = play.option_order(6, True, "solo:1:0")
    assert sorted(shuffled) == list(range(6)) and shuffled == play.option_order(6, True, "solo:1:0")
    assert play.points_for(True, 0, 20, 3, speed=False, streak_bonus=False) == 1000


def test_cards_render_every_screen():
    from studybot import cards

    q = {"stem": "Hồ Chí Minh đã tiếp thu những giá trị tích cực nào của Phật giáo? 🙏",
         "options": ["Tư tưởng từ bi", "Bình đẳng", "Vô vi", "Lễ giáo", "x" * 400], "correct_index": 2}
    for png in (cards.question_card(q, 1, 10, right="1,200 pts · #1/3", timer="20s"),
                cards.reveal_card(q, 1, 10, order=[4, 3, 2, 1, 0], chosen=0, counts=[1, 0, 5, 0, 2], banner="Answer: C"),
                cards.podium_card([("Ada 🌸", 3000), ("Bo", 2000)]), cards.score_card(7, 10)):
        assert png[:8] == b"\x89PNG\r\n\x1a\n" and len(png) < 1_500_000
    assert cards.clean("Lan Anh 🌸✨") == "Lan Anh"


def test_cards_show_long_answers_whole():
    from studybot import cards

    long = "Tư tưởng Hồ Chí Minh là kết quả của sự vận dụng sáng tạo và phát triển chủ nghĩa Mác - Lênin vào điều kiện"
    q = {"stem": "Chọn phương án trả lời đúng nhất?", "options": [long, long + " cụ thể", "Ngắn", long], "correct_index": 0}
    layout = cards.plan(q["stem"], q["options"])
    assert (layout.cols, layout.rows) == (2, 2) and layout.text_size >= 34 and layout.height <= cards.GROW_H
    tile_w = (cards.W - 2 * cards.PAD - cards.GAP) // 2 - 44
    for o in q["options"]:  # every answer fits whole at the chosen size: nothing is cut with "…"
        assert cards.text_height(o, tile_w, layout.text_size) + cards._tile_reserved(None) <= layout.tile_h
    short = cards.plan("Q?", ["Đại hội VII", "Đại hội IX", "Đại hội XI", "Đại hội XIII"])
    assert (short.cols, short.rows, short.height) == (4, 1, cards.H)  # short answers keep the one-row card


def test_solo_run_with_timer_and_hidden_answers(tmp_path):
    import types
    from studybot import play

    store, deck = _deck_with_questions(2)
    ids = [r["id"] for r in store.deck_questions(deck)]
    settings = play.QuizSettings(timer=10, show_answers=False, shuffle=True, cards=False)
    run_id = store.create_run(7, ids, settings.to_dict())
    log = []

    async def edit_message(**kw):
        log.append(kw)

    inter = types.SimpleNamespace(client=types.SimpleNamespace(store=store), user=types.SimpleNamespace(id=7),
                                  response=types.SimpleNamespace(edit_message=edit_message, send_message=edit_message))

    async def scenario():
        run = store.get_run(run_id)
        await play.solo_question(store, run)
        q = store.get_question(ids[0])
        shown = play.run_order(run, q, 0).index(q["correct_index"])
        await play.SoloAnswer(run_id, 0, shown).callback(inter)  # right answer, on time, shuffled tile
        assert "Locked in" in json.dumps(log[-1]["view"].to_components()) and log[-1]["attachments"] == []
        run = store.get_run(run_id)
        await play.solo_question(store, run)
        store.mark_asked(run_id, time.time() - 60)  # pretend the timer ran out
        await play.SoloAnswer(run_id, 1, 0).callback(inter)
        assert "Too slow" in json.dumps(log[-1]["view"].to_components())
        run = store.get_run(run_id)
        assert run["correct"] == 1 and run["picks"] == [q["correct_index"], None]
        view, files = await play.solo_summary(store, run)
        text = json.dumps(view.to_components(), ensure_ascii=False)
        assert "1/2 correct" in text and "⏰ no answer" in text and files == []

    asyncio.run(scenario())


def test_live_shuffle_scoring_and_personal_results():
    import types
    from studybot import play

    store, deck = _deck_with_questions(1)
    qid = store.deck_questions(deck)[0]["id"]
    q = store.get_question(qid)
    sent, edits = [], []

    class Msg:
        embeds = []

        async def edit(self, **kw):
            edits.append(kw)

    class Channel:
        async def send(self, *a, **kw):
            sent.append(kw)
            return Msg()

    client = types.SimpleNamespace(store=store, live_games={})
    replies = []

    def inter(uid):
        async def send_message(text, **kw):
            replies.append(text)

        async def delete_original_response():
            pass

        return types.SimpleNamespace(client=client, user=types.SimpleNamespace(id=uid, display_name=f"p{uid}"),
                                     response=types.SimpleNamespace(send_message=send_message),
                                     delete_original_response=delete_original_response)

    async def scenario():
        s = play.QuizSettings(shuffle=True, speed=False, streak=False, cards=True)
        game = play.LiveGame(5, host_id=1, channel=Channel(), question_ids=[qid], time_limit=30, settings=s)
        client.live_games[5] = game
        await play.ask_next(client, game)
        # One picture: the card as a plain attachment, no embed pointing at it (Discord would draw it twice).
        assert sent[-1]["files"] and sent[-1]["embed"] is None and "**0** answered" in sent[-1]["content"]
        right = game.order.index(q["correct_index"])
        await play.LiveAnswer(5, 0, right).callback(inter(2))
        await play.LiveAnswer(5, 0, (right + 1) % len(game.order)).callback(inter(3))
        assert game.scores == {2: 1000, 3: 0} and "#1 of 1" in replies[0] and "#2 of 2" in replies[1]
        await asyncio.sleep(play.COUNTER_EDIT_GAP + 0.1)
        counter = edits[-1]
        assert "**2** answered" in counter["content"] and "attachments" not in counter and counter["embed"] is None
        await play.reveal(game, 0)
        assert edits[-1]["attachments"]  # reveal card replaces the question card on the same message
        assert not edits[-1]["embed"].image  # results text under the card, not a second copy of it
        await play.LiveMine(5, 0).callback(inter(3))
        assert "no points" in replies[-1] and "#2 of 2" in replies[-1]
        await play.LiveMine(5, 0).callback(inter(4))
        assert "didn't answer" in replies[-1]
        await play.ask_next(client, game)
        assert "🥇 **p2**" in edits[-1]["embed"].description and edits[-1]["attachments"]

    asyncio.run(scenario())


def test_live_auto_advance_and_host_next_cancels_it():
    import types
    from studybot import play

    store, _ = _deck_with_questions(3)
    ids = [q["id"] for q in store.deck_questions(store.find_deck(9, "enzymes 101"))][:3]
    edits = []

    class Msg:
        async def edit(self, **kw):
            edits.append(kw)

    class Channel:
        async def send(self, *a, **kw):
            return Msg()

    client = types.SimpleNamespace(store=store, live_games={})

    async def scenario():
        s = play.QuizSettings(cards=False, auto_next=0.05)  # real choices are 3s+; short for the test
        game = play.LiveGame(7, host_id=1, channel=Channel(), question_ids=ids, time_limit=30, settings=s,
                             client=client)
        client.live_games[7] = game
        await play.ask_next(client, game)
        await play.reveal(game, 0)
        assert "Next question <t:" in edits[-1]["content"]
        await asyncio.sleep(0.2)
        assert game.index == 1 and not game.revealed  # moved on by itself

        await play.reveal(game, 1)
        pending = game.timer
        await play.ask_next(client, game)  # host pressed Next first
        await asyncio.sleep(0)
        assert pending.cancelled()
        assert game.index == 2
        await play.reveal(game, 2)
        assert "Final results <t:" in edits[-1]["content"]
        await asyncio.sleep(0.2)
        assert 7 not in client.live_games and "Game over" in edits[-1]["embed"].title
        game.timer and game.timer.cancel()

    asyncio.run(scenario())
    assert play.QuizSettings.from_dict({"auto_next": 7}).auto_next == 0  # not a listed choice
    assert play.QuizSettings.from_dict({"auto_next": 5}).auto_next == 5


def test_powerup_rules():
    import random as _random
    from studybot import powerups as pu

    rng = _random.Random(1)
    p = pu.Powers()
    assert pu.earned_after_answer(p, False, 0, rng) in pu.CATALOGUE  # starter, even on a wrong answer
    assert pu.earned_after_answer(p, True, 1, rng) is None
    assert pu.earned_after_answer(p, True, 3, rng) and len(p.held) == 2
    p.held += ["x", "y"]
    assert pu.award(p, rng) is None  # full hands
    p = pu.Powers()
    assert pu.erase(p, 4, correct=2, count=2, rng=rng) and len(p.removed) == 2 and 2 not in p.removed
    assert pu.erase(p, 4, correct=2, count=1, rng=rng) == []  # always leaves one wrong answer
    p.active = {"double"}
    assert pu.adjust(800, True, p, power_play=True, right_value=800)[0] == 2400
    p.active = {"jeopardy"}
    assert pu.adjust(0, False, p, power_play=False, right_value=900)[0] == -900
    p.active, p.supersonic = set(), 1
    assert pu.adjust(1000, True, p, False, 1000)[0] == 1500 and p.supersonic == 0
    p.frozen_left = 18.0
    assert pu.effective_left(p, 2.0) == 18.0
    p.held = ["double", "jeopardy", "gift"]
    p.active = {"double"}
    assert "don't stack" in pu.refuse_reason("jeopardy", p, False, False, 3, 1)
    assert "before you pick" in pu.refuse_reason("double", pu.Powers(held=["double"]), True, False, 3, 1)
    assert "Nobody else" in pu.refuse_reason("gift", p, False, False, 3, 0)
    assert pu.refuse_reason("gift", p, True, False, 3, 2) is None  # gifts work after answering


def test_live_powerups_in_a_game():
    import types
    from studybot import play, powerups as pu

    store, _ = _deck_with_questions(2)
    ids = [q["id"] for q in store.deck_questions(store.find_deck(9, "enzymes 101"))][:2]
    edits = []

    class Msg:
        async def edit(self, **kw):
            edits.append(kw)

    class Channel:
        async def send(self, *a, **kw):
            return Msg()

    client = types.SimpleNamespace(store=store, live_games={})

    def inter(uid):
        log = []

        async def send_message(text=None, **kw):
            log.append((text or kw.get("content"), kw))

        async def edit_message(**kw):
            log.append((kw.get("content"), kw))

        async def delete_original_response():
            pass

        return types.SimpleNamespace(client=client, user=types.SimpleNamespace(id=uid, display_name=f"p{uid}"), log=log,
                                     response=types.SimpleNamespace(send_message=send_message, edit_message=edit_message),
                                     delete_original_response=delete_original_response)

    async def scenario():
        s = play.QuizSettings(cards=False, speed=False, streak=False)
        game = play.LiveGame(9, host_id=1, channel=Channel(), question_ids=ids, time_limit=30, settings=s, client=client)
        client.live_games[9] = game
        await play.ask_next(client, game)
        right = game.shown_correct
        wrong = [i for i in range(len(game.order)) if i != right]

        # p2 holds 50-50, Immunity and 2x; p3 holds Double Jeopardy; p4 holds Gift and Power Play.
        game.powers[2] = pu.Powers(held=["fifty", "immunity", "double"], started=True)
        game.powers[3] = pu.Powers(held=["jeopardy"], started=True)
        game.powers[4] = pu.Powers(held=["gift", "power_play"], started=True)
        menu = inter(2)
        await play.LivePowerMenu(9, 0).callback(menu)
        assert "50-50" in menu.log[0][0] and menu.log[0][1]["ephemeral"]
        await play.LivePowerUse(9, 0, "fifty").callback(menu)
        assert len(game.powers[2].removed) == 2 and "gone" in menu.log[-1][0]
        erased = sorted(game.powers[2].removed)[0]
        tap = inter(2)
        await play.LiveAnswer(9, 0, erased).callback(tap)
        assert "erased for you" in tap.log[0][0] and 2 not in game.answers
        await play.LivePowerUse(9, 0, "immunity").callback(menu)
        await play.LivePowerUse(9, 0, "double").callback(menu)
        last_wrong = next(i for i in wrong if i not in game.powers[2].removed)
        await play.LiveAnswer(9, 0, last_wrong).callback(tap)
        assert "Immunity" in tap.log[-1][0] and 2 not in game.answers  # second try
        await play.LiveAnswer(9, 0, right).callback(tap)
        assert game.scores[2] == 2000 and "2x" in tap.log[-1][0]

        p3 = inter(3)
        await play.LivePowerUse(9, 0, "jeopardy").callback(p3)
        await play.LiveAnswer(9, 0, wrong[0]).callback(p3)
        assert game.scores[3] == 0 and "🎲" in p3.log[-1][0]  # can't go below zero

        p4 = inter(4)
        await play.LiveAnswer(9, 0, right).callback(p4)
        assert game.scores[4] == 1000
        await play.LivePowerUse(9, 0, "power_play").callback(p4)
        assert game.scores[4] == 1500 and game.scores[2] == 3000  # right answers already in get +50% too
        before = dict(game.scores)
        await play.LivePowerUse(9, 0, "gift").callback(p4)
        gifted = [u for u in (2, 3) if game.scores[u] == before[u] + pu.GIFT_POINTS]
        assert len(gifted) == 1 and any("gifted" in e for e in game.events) and any("Power Play" in e for e in game.events)
        await asyncio.sleep(play.COUNTER_EDIT_GAP + 0.1)
        assert "gifted" in edits[-1]["content"]  # news shows on the public message
        await play.reveal(game, 0)
        await play.ask_next(client, game)
        assert not game.powers[2].removed and not game.powers[2].active and not game.events
        game.timer and game.timer.cancel()

    asyncio.run(scenario())


def test_music_tracks_and_cues(tmp_path):
    import types
    import wave
    from studybot import music

    paths = music.ensure_tracks(tmp_path)
    for name, path in paths.items():
        with wave.open(str(path)) as w:
            assert w.getframerate() == music.SR and w.getnframes() > music.SR  # at least a second of audio
    played, afters = [], []

    class FakeVC:
        client = types.SimpleNamespace(loop=types.SimpleNamespace(call_soon_threadsafe=lambda fn: fn()))

        def is_connected(self):
            return True

        def is_playing(self):
            return bool(played)

        def is_paused(self):
            return False

        def stop(self):
            pass

        def play(self, source, after=None):
            played.append(source)
            afters.append(after)

    assert len(music.QUESTION_TRACKS) >= 6 and len(music.LOBBY_TRACKS) >= 2
    (tmp_path / "stale-v0.wav").write_bytes(b"x")
    assert not (music.ensure_tracks(tmp_path) and (tmp_path / "stale-v0.wav").exists())  # old versions cleaned up
    gm = music.GameMusic()
    gm.vc, gm.paths = FakeVC(), paths
    gm.rotation = [paths[k] for k in music.QUESTION_TRACKS]
    picks = [gm.pick("question") for _ in range(len(gm.rotation) + 1)]
    assert len(set(picks)) == len(gm.rotation) and picks[-1] == picks[0]  # every loop once, then round again
    assert gm.pick("lobby") != gm.pick("lobby")
    gm.cue("reveal", then="lobby")
    gm.cue("question")  # host skipped ahead during the sting
    afters[0](None)     # the stopped sting finishing must not start the lobby loop over the question music
    assert len(played) == 2
    gm.cue("reveal", then="lobby")
    afters[-1](None)
    assert len(played) == 4  # sting finished normally → lobby loop


def test_quiz_panel_fits_discord_limits():
    from studybot import play

    store, deck = _deck_with_questions(3)
    for i in range(30):
        store.get_or_create_deck(9, f"extra set {i}", subject=f"subject {i % 30}")
    picker = play.QuizPicker(store, 9, owner=1)

    def count(items):
        n = 0
        for c in items:
            n += 1
            for key in ("components",):
                n += count(c.get(key, []))
            if "accessory" in c:
                n += 1
        return n

    payload = picker.to_components()
    assert count(payload) <= 40
    texts = json.dumps(payload)
    assert "Shuffle answer choices" in texts and "Quizizz-style cards" in texts
    store.set_prefs(1, {"length": 25, "shuffle": True})
    again = play.QuizPicker(store, 9, owner=1)
    assert again.settings.length == 25 and again.settings.shuffle


def _count_components(items):
    return sum(1 + _count_components(c.get("components", [])) + (1 if "accessory" in c else 0) for c in items)


def test_library_organises_subjects_and_adds_files():
    import types
    from studybot import library
    from studybot.bot import StudyBot, StudyCog

    async def scenario():
        bot = StudyBot(Store(":memory:"))
        cog = StudyCog(bot)
        store = bot.store
        for i in range(14):
            store.get_or_create_deck(9, f"set {i}", subject=f"subject {i}")
        panel = library.LibraryPanel(cog, 9, owner=1)
        assert _count_components(panel.to_components()) <= 40

        bank = json.dumps([{"q": f"Question {i}?", "options": ["a", "b", "c"], "correctIndex": 1} for i in range(3)]).encode()
        sent = []

        async def defer(**kw):
            pass

        async def followup(text, **kw):
            sent.append(text)

        class Att:
            def __init__(self, name, data):
                self.filename, self._data = name, data

            async def read(self):
                return self._data

        # A brand-new subject, one new set named in the form, two files.
        modal = library.AddFilesModal(cog, 9, None)
        modal.subject_input._value = "Tư tưởng HCM"
        modal.files._values = [Att("week1.json", bank), Att("week2.json", bank.replace(b"Question", b"Cau"))]
        modal.target._values = [library.NEW_NAMED]
        modal.set_name._value = "Review"
        inter = types.SimpleNamespace(user=types.SimpleNamespace(id=1),
                                      response=types.SimpleNamespace(defer=defer),
                                      followup=types.SimpleNamespace(send=followup))
        await modal.on_submit(inter)
        deck = store.find_deck(9, "review")
        assert deck and store.deck_subject(deck) == "tư tưởng hcm" and len(store.deck_questions(deck)) == 6
        assert all("Found **3**" in t for t in sent)

        # Move and rename it from its set panel.
        edits = []

        async def edit_message(**kw):
            edits.append(kw)

        inter2 = types.SimpleNamespace(user=types.SimpleNamespace(id=1), response=types.SimpleNamespace(edit_message=edit_message))
        set_panel = library.SetPanel(cog, 9, 1, deck)
        await set_panel._apply_subject(inter2, "history")
        assert store.deck_subject(deck) == "history"
        assert store.rename_deck(deck, "Week reviews") and store.find_deck(9, "week reviews") == deck
        assert not store.rename_deck(deck, "set 1")  # name taken
        await bot.close()

    asyncio.run(scenario())


def test_live_game_checks_channel_permissions():
    from studybot import play

    perms = discord.Permissions(view_channel=True, send_messages=True)
    assert play.missing_live_permissions(perms) == ["Embed Links", "Attach Files", "Read Message History"]
    assert play.missing_live_permissions(discord.Permissions.all()) == []
    thread_perms = discord.Permissions(view_channel=True, send_messages_in_threads=True, embed_links=True,
                                       attach_files=True, read_message_history=True)
    assert play.missing_live_permissions(thread_perms, in_thread=True) == []
    assert play.missing_live_permissions(thread_perms) == ["Send Messages"]
    assert "Edit Channel" in play.permission_help(["View Channel"])
