"""Turns uploaded files into clean text chunks the question writer can cite from."""
import hashlib
import io
import json
import re
from pathlib import PurePath

from . import config

TEXT_EXTENSIONS = {".txt", ".md", ".markdown", ".rst", ".csv", ".html", ".htm"}
# Read by Gemini directly (photos of worksheets, screenshots), since there's no text layer.
IMAGE_TYPES = {".png": "image/png", ".jpg": "image/jpeg", ".jpeg": "image/jpeg", ".webp": "image/webp",
               ".heic": "image/heic", ".heif": "image/heif"}
SUPPORTED_EXTENSIONS = TEXT_EXTENSIONS | set(IMAGE_TYPES) | {".pdf", ".docx", ".pptx", ".json"}
OPTION_LIKE = re.compile(r"^\s*[\(\[]?[A-Fa-f][\)\]\.:]\s+\S")


class UnsupportedFile(ValueError):
    pass


# Limits for untrusted uploads: a 25 MB .docx/.pptx can unzip to gigabytes (a "zip bomb"),
# and a PDF can declare thousands of pages.
MAX_UNZIPPED_BYTES = 200 * 1024 * 1024
MAX_ZIP_ENTRIES = 5000
MAX_PDF_PAGES = 500


def check_zip(data: bytes) -> None:
    import zipfile

    try:
        with zipfile.ZipFile(io.BytesIO(data)) as z:
            infos = z.infolist()
    except zipfile.BadZipFile as exc:
        raise UnsupportedFile("That file looks damaged (not a valid Office file).") from exc
    if len(infos) > MAX_ZIP_ENTRIES or sum(i.file_size for i in infos) > MAX_UNZIPPED_BYTES:
        raise UnsupportedFile("That file unpacks to something far too large to read safely.")


def extract_text(filename: str, data: bytes) -> str:
    ext = PurePath(filename).suffix.lower()
    if ext in {".docx", ".pptx"}:
        check_zip(data)
    if ext == ".pdf":
        from pypdf import PdfReader

        reader = PdfReader(io.BytesIO(data))
        if len(reader.pages) > MAX_PDF_PAGES:
            raise UnsupportedFile(f"That PDF has {len(reader.pages)} pages; the limit is {MAX_PDF_PAGES}.")
        pages = [page.extract_text() or "" for page in reader.pages]
        text = "\n\n".join(pages)
    elif ext == ".docx":
        text = _docx_text(data)
    elif ext == ".pptx":
        from pptx import Presentation

        slides = []
        for slide in Presentation(io.BytesIO(data)).slides:
            parts = []
            for shape in slide.shapes:
                if shape.has_text_frame:
                    parts.append(shape.text_frame.text)
                elif getattr(shape, "has_table", False) and shape.has_table:
                    parts += [cell.text for row in shape.table.rows for cell in row.cells]
            slides.append("\n".join(parts))
        text = "\n\n".join(slides)
    elif ext == ".json":
        text = json_text(data)
    elif ext in IMAGE_TYPES:
        return ""
    elif ext in TEXT_EXTENSIONS:
        text = data.decode("utf-8", errors="replace")
        if ext in {".html", ".htm"}:
            text = re.sub(r"<(script|style)\b.*?</\1>", " ", text, flags=re.S | re.I)
            text = re.sub(r"<[^>]+>", " ", text)
    else:
        raise UnsupportedFile(
            f"Can't read `{ext or filename}` files. Supported: {', '.join(sorted(SUPPORTED_EXTENSIONS))}"
        )
    return normalize(text)


def json_text(data: bytes) -> str:
    """Every string value in a JSON file, one per line, so notes and question banks read as text.
    Invalid JSON is read as plain text."""
    raw = data.decode("utf-8-sig", errors="replace")
    try:
        root = json.loads(raw)
    except (json.JSONDecodeError, RecursionError):
        return raw
    out, stack = [], [root]
    while stack:  # iterative, so deeply nested files can't overflow the stack
        node = stack.pop()
        if isinstance(node, dict):
            stack.extend(reversed(list(node.values())))
        elif isinstance(node, list):
            stack.extend(reversed(node))
        elif isinstance(node, str) and node.strip():
            out.append(node.strip())
    return "\n".join(out)


def _docx_text(data: bytes) -> str:
    """Paragraphs and table cells. An option set in bold or highlight (a common way
    teachers mark the answer) gets a trailing "*" so the parser can read the key."""
    import docx

    document = docx.Document(io.BytesIO(data))

    def para_text(p) -> str:
        text = p.text
        runs = [r for r in p.runs if r.text.strip()]
        if OPTION_LIKE.match(text) and runs and (
            all(r.bold for r in runs) or any(r.font.highlight_color for r in runs)
        ):
            text += " *"
        return text

    lines = [para_text(p) for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                lines += [para_text(p) for p in cell.paragraphs]
            lines.append("")
    return "\n".join(lines)


def needs_vision(filename: str, text: str) -> bool:
    """Images, and PDFs with no usable text layer (scans), go to the model as files."""
    ext = PurePath(filename).suffix.lower()
    return ext in IMAGE_TYPES or (ext == ".pdf" and len(text) < 200)


def mime_type(filename: str) -> str:
    ext = PurePath(filename).suffix.lower()
    return IMAGE_TYPES.get(ext, "application/pdf")


def normalize(text: str) -> str:
    text = text.replace("\r\n", "\n").replace("\x00", "")
    # Re-join words hyphenated across PDF line breaks.
    text = re.sub(r"(\w)-\n(\w)", r"\1\2", text)
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def chunk_text(text: str, size: int = config.CHUNK_CHARS, overlap: int = config.CHUNK_OVERLAP) -> list[str]:
    """Splits on paragraph boundaries into ~size-char chunks with a little overlap,
    so a question's supporting passage is rarely cut in half."""
    paragraphs = [p.strip() for p in text.split("\n\n") if p.strip()]
    chunks: list[str] = []
    current = ""
    for para in paragraphs:
        while len(para) > size:
            # One huge paragraph (common in PDFs): break it at a sentence end.
            cut = para.rfind(". ", 0, size)
            cut = cut + 1 if cut > size // 2 else size
            pieces = para[:cut].strip(), para[cut:].strip()
            if current:
                chunks.append(current)
                current = ""
            chunks.append(pieces[0])
            para = pieces[1]
        if current and len(current) + len(para) + 2 > size:
            chunks.append(current)
            tail = current[-overlap:]
            space = tail.find(" ")
            current = (tail[space + 1:] if space != -1 else tail) + "\n\n" + para
        else:
            current = f"{current}\n\n{para}" if current else para
    if current:
        chunks.append(current)
    # Fragments too small to write a question from are only noise.
    return [c for c in chunks if len(c) >= 200] or chunks


def sha256(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()
